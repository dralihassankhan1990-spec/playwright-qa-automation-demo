#!/usr/bin/env python3
"""
Excel Task Tracker Builder
==========================
Builds a formatted Excel work tracker (.xlsx) from a JSON task list.

Features:
  - Styled header row, frozen panes, auto-filters
  - Status column with a dropdown (Not Started / In Progress / Blocked / Done)
  - RAG conditional formatting on Status and Priority
  - Summary sheet with live COUNTIF formulas

Usage:
    python excel_tracker.py sample_tasks.json
    python excel_tracker.py tasks.json -o tracker_oct.xlsx

JSON schema:
{
  "title": "Q4 Work Tracker",
  "tasks": [
    {"id": "T-01", "task": "...", "owner": "...",
     "status": "In Progress", "priority": "High", "due": "2026-10-20"},
    ...
  ]
}
"""
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ["ID", "Task", "Owner", "Status", "Priority", "Due Date"]
STATUSES = ["Not Started", "In Progress", "Blocked", "Done"]
PRIORITIES = ["Low", "Medium", "High", "Critical"]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build(data: dict, out_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"
    tasks = data.get("tasks", [])

    # Title + header
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = data.get("title", "Work Tracker")
    c.font = TITLE_FONT
    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.row_dimensions[3].height = 22

    # Rows
    for i, t in enumerate(tasks, start=4):
        vals = [t.get("id", ""), t.get("task", ""), t.get("owner", ""),
                t.get("status", ""), t.get("priority", ""), t.get("due", "")]
        for col, v in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=col, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 2))
        ws.row_dimensions[i].height = 20

    last = 3 + len(tasks)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 14
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:F{last}"

    # Status dropdown
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D4:D{last}")

    # RAG formatting on Status
    colors = {"Done": "C6EFCE", "In Progress": "FFEB9C",
              "Blocked": "FFC7CE", "Not Started": "D9D9D9"}
    for status, color in colors.items():
        ws.conditional_formatting.add(
            f"D4:D{last}",
            FormulaRule(formula=[f'$D4="{status}"'],
                        fill=PatternFill("solid", fgColor=color)))

    # Priority formatting
    ws.conditional_formatting.add(
        f"E4:E{last}",
        CellIsRule(operator="equal", formula=['"Critical"'],
                   fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add(
        f"E4:E{last}",
        CellIsRule(operator="equal", formula=['"High"'],
                   fill=PatternFill("solid", fgColor="FFEB9C")))

    # Summary sheet
    sm = wb.create_sheet("Summary")
    sm["A1"] = data.get("title", "Work Tracker") + " — Summary"
    sm["A1"].font = TITLE_FONT
    sm["A3"] = "Metric"
    sm["B3"] = "Count"
    for cell in (sm["A3"], sm["B3"]):
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    metrics = [("Total tasks", f"=COUNTA(Tasks!A4:A{last})")]
    metrics += [(s, f'=COUNTIF(Tasks!D4:D{last},"{s}")') for s in STATUSES]
    metrics.append(("Overdue (not Done)",
                    f'=COUNTIFS(Tasks!F4:F{last},"<"&TODAY(),Tasks!D4:D{last},"<>Done")'))
    for i, (label, formula) in enumerate(metrics, start=4):
        sm.cell(row=i, column=1, value=label).font = Font(bold=True)
        sm.cell(row=i, column=2, value=formula)
    sm.column_dimensions["A"].width = 24
    sm.column_dimensions["B"].width = 12

    wb.save(out_path)
    print(f"Saved: {out_path} ({len(tasks)} tasks)")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    out = sys.argv[3] if len(sys.argv) > 3 and sys.argv[2] == "-o" else "work_tracker.xlsx"
    with open(src, encoding="utf-8") as f:
        build(json.load(f), out)


if __name__ == "__main__":
    main()
