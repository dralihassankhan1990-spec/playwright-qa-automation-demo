#!/usr/bin/env python3
"""
CSV Consolidator
================
Merges every CSV file in a folder into a single Excel workbook (.xlsx):
one worksheet per CSV plus a Summary sheet listing each file's row
and column counts.

Usage:
    python csv_consolidator.py ./csv_files
    python csv_consolidator.py ./csv_files -o consolidated.xlsx

Each CSV's first row is treated as the header. Sheet names come from
the file names (sanitized for Excel's 31-character / character rules).
"""
import csv
import os
import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sheet_name(filename: str, used: set) -> str:
    base = os.path.splitext(os.path.basename(filename))[0]
    name = re.sub(r'[\\/*?:\[\]]', "_", base)[:31] or "Sheet"
    i, candidate = 2, name
    while candidate in used:
        candidate = f"{name[:28]}_{i}"
        i += 1
    used.add(candidate)
    return candidate


def consolidate(folder: str, out_path: str):
    files = sorted(f for f in os.listdir(folder) if f.lower().endswith(".csv"))
    if not files:
        print(f"No CSV files found in {folder}")
        sys.exit(1)

    wb = Workbook()
    first = True
    used, stats = set(), []
    for fname in files:
        path = os.path.join(folder, fname)
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
        if not rows:
            continue
        header, data = rows[0], rows[1:]
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = sheet_name(fname, used)
        for col, h in enumerate(header, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
        for r, row in enumerate(data, start=2):
            for col, v in enumerate(row, start=1):
                cell = ws.cell(row=r, column=col, value=v)
                cell.border = BORDER
        for col in range(1, len(header) + 1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 20
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        stats.append((fname, len(data), len(header)))

    sm = wb.create_sheet("Summary", 0)
    sm["A1"] = "Consolidation Summary"
    sm["A1"].font = TITLE_FONT
    for col, h in enumerate(["Source File", "Rows", "Columns"], start=1):
        cell = sm.cell(row=3, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
    for i, (fname, nrows, ncols) in enumerate(stats, start=4):
        sm.cell(row=i, column=1, value=fname).border = BORDER
        sm.cell(row=i, column=2, value=nrows).border = BORDER
        sm.cell(row=i, column=3, value=ncols).border = BORDER
    sm.cell(row=4 + len(stats), column=1, value="TOTAL").font = Font(bold=True)
    sm.cell(row=4 + len(stats), column=2,
            value=f"=SUM(B4:B{3 + len(stats)})").font = Font(bold=True)
    for col in ("A", "B", "C"):
        sm.column_dimensions[col].width = 28 if col == "A" else 12

    wb.save(out_path)
    print(f"Saved: {out_path} ({len(stats)} CSVs consolidated)")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    folder = sys.argv[1]
    out = sys.argv[3] if len(sys.argv) > 3 and sys.argv[2] == "-o" else "consolidated.xlsx"
    consolidate(folder, out)


if __name__ == "__main__":
    main()
