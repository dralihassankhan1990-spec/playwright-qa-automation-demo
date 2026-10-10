"""
Excel-driven Playwright automation demo (portfolio project).

Reads test cases from test_data.xlsx, runs each one in Chromium,
writes PASS/FAIL back to output/results.xlsx (plus extra report sheets),
and saves a screenshot + video per case.

Run:  python demo.py
Requires: pip install -r requirements.txt && playwright install chromium
"""
import os
import re
import sys
import traceback
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from playwright.sync_api import sync_playwright

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_XLSX = os.path.join(HERE, "test_data.xlsx")
OUT_DIR = os.path.join(HERE, "output")
SHOT_DIR = os.path.join(OUT_DIR, "screenshots")
VIDEO_DIR = os.path.join(OUT_DIR, "videos")
RESULTS_XLSX = os.path.join(OUT_DIR, "results.xlsx")

BASE = "https://www.saucedemo.com"
SPRINT_URL = "https://todomvc.com/examples/react/dist/"

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
GREEN_FONT = Font(color="006100", bold=True)
RED_FONT = Font(color="9C0006", bold=True)
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Border(*[Side(style="thin", color="B0B0B0")] * 4)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

STEP_PAUSE = 700  # ms between steps so the run is easy to follow on video


def _proxy_or_none():
    """Use the bundled forward proxy if it is running (restricted networks)."""
    import socket
    try:
        s = socket.create_connection(("127.0.0.1", 8888), timeout=1)
        s.close()
        return {"server": "http://127.0.0.1:8888"}
    except OSError:
        return None


def highlight(page, locator, ms=400):
    """Brief red outline on the element being acted on (video clarity)."""
    try:
        locator.evaluate(
            "e => { e.style.outline='3px solid #e11d48'; e.style.outlineOffset='2px'; }")
        page.wait_for_timeout(ms)
        locator.evaluate("e => { e.style.outline=''; e.style.outlineOffset=''; }")
    except Exception:
        pass


def load_cases():
    wb = load_workbook(DATA_XLSX)
    ws = wb.active
    cases = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        cases.append({
            "name": str(row[0]).strip(),
            "username": str(row[1]).strip() if row[1] else "",
            "password": str(row[2]).strip() if row[2] else "",
            "products": [p.strip() for p in str(row[3]).split(",")] if row[3] else [],
            "expected": str(row[4]).strip().lower() if row[4] else "pass",
        })
    return cases


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN


def new_results_wb():
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["Case", "Status", "Notes", "Run at"])
    style_header(ws, 4)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 20
    return wb


def write_kpi_sheet(wb, items):
    ws = wb.create_sheet("Catalog_KPI_Report")
    ws.append(["Product", "Price (USD)"])
    style_header(ws, 2)
    prices = []
    for it in items:
        ws.append([it["name"], it["price"]])
        prices.append(it["price"])
    n = len(prices)
    avg = round(sum(prices) / n, 2) if n else 0
    ws.append([])
    ws.append(["Total products", n])
    ws.append(["Average price", avg])
    ws.append(["Highest price", max(prices) if prices else 0])
    ws.append(["Lowest price", min(prices) if prices else 0])
    for r in range(2, ws.max_row + 1):
        for c in range(1, 3):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN
            cell.alignment = LEFT if c == 1 else CENTER
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    return {"total": n, "avg": avg}


def write_sprint_sheet(wb, tasks):
    ws = wb.create_sheet("Sprint_Summary")
    ws.append(["Sprint task", "Status"])
    style_header(ws, 2)
    for name, done in tasks:
        ws.append([name, "Done" if done else "To Do"])
    for r in range(2, ws.max_row + 1):
        for c in range(1, 3):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN
            cell.alignment = LEFT if c == 1 else CENTER
        status = ws.cell(row=r, column=2)
        if status.value == "Done":
            status.fill = GREEN
            status.font = GREEN_FONT
        else:
            status.fill = PatternFill("solid", fgColor="FFEB9C")
            status.font = Font(color="9C6500", bold=True)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 14
    done = sum(1 for _, d in tasks if d)
    return {"done": done, "total": len(tasks)}


# ---------------------------------------------------------------- flows

def do_login(page, username, password):
    page.goto(BASE, wait_until="domcontentloaded")
    page.wait_for_selector("#user-name", timeout=15000)
    page.wait_for_timeout(STEP_PAUSE)
    u = page.locator("#user-name")
    highlight(page, u); u.fill(username)
    page.wait_for_timeout(STEP_PAUSE // 2)
    p = page.locator("#password")
    highlight(page, p); p.fill(password)
    page.wait_for_timeout(STEP_PAUSE // 2)
    b = page.locator("#login-button")
    highlight(page, b); b.click()


def do_purchase(page, case):
    """expected == 'pass': log in, add products, verify cart badge."""
    do_login(page, case["username"], case["password"])
    page.wait_for_selector(".inventory_item", timeout=15000)
    page.wait_for_timeout(STEP_PAUSE)
    added = 0
    for product in case["products"]:
        card = page.locator(".inventory_item", has_text=product).first
        card.scroll_into_view_if_needed()
        btn = card.get_by_role("button", name=re.compile("add to cart", re.I))
        highlight(page, btn); btn.click()
        added += 1
        page.wait_for_timeout(STEP_PAUSE)
    badge = page.locator(".shopping_cart_badge")
    badge.wait_for(timeout=8000)
    assert badge.inner_text().strip() == str(added), \
        f"cart badge shows {badge.inner_text()!r}, expected {added}"
    cart = page.locator(".shopping_cart_link")
    highlight(page, cart); cart.click()
    page.wait_for_selector(".cart_item", timeout=8000)
    page.wait_for_timeout(STEP_PAUSE)
    return f"Logged in, added {added} product(s), cart verified"


def do_locked_out(page, case):
    """expected == 'fail': locked-out user must be rejected at login."""
    do_login(page, case["username"], case["password"])
    err = page.locator("[data-test='error']")
    err.wait_for(timeout=8000)
    page.wait_for_timeout(STEP_PAUSE)
    assert "locked out" in err.inner_text().lower(), "unexpected error text"
    return "Login correctly rejected (locked out user)"


def do_report(page, case, wb):
    """expected == 'report': scrape catalog, generate KPI report sheet."""
    do_login(page, case["username"], case["password"])
    page.wait_for_selector(".inventory_item", timeout=15000)
    page.wait_for_timeout(STEP_PAUSE)
    # scroll through the catalog so the extraction is visible
    for i in range(3):
        page.mouse.wheel(0, 400)
        page.wait_for_timeout(400)
    items = page.eval_on_selector_all(
        ".inventory_item",
        "els => els.map(e => ({"
        " name: e.querySelector('.inventory_item_name').innerText,"
        " price: parseFloat(e.querySelector('.inventory_item_price')"
        "   .innerText.replace('$','')) }))")
    assert len(items) >= 6, f"expected 6 catalog products, found {len(items)}"
    stats = write_kpi_sheet(wb, items)
    page.wait_for_timeout(STEP_PAUSE)
    return (f"Extracted {stats['total']} products to Catalog_KPI_Report — "
            f"avg price ${stats['avg']}")


def do_sprint(page, case, wb):
    """expected == 'sprint': drive a sprint board, summarize progress."""
    page.goto(SPRINT_URL, wait_until="domcontentloaded")
    box = page.get_by_placeholder("What needs to be done?")
    box.wait_for(timeout=15000)
    page.wait_for_timeout(STEP_PAUSE)
    tasks = ["Draft sprint plan", "Review UAT results", "Prepare go-live checklist"]
    for t in tasks:
        highlight(page, box)
        box.fill(t)
        box.press("Enter")
        page.wait_for_timeout(STEP_PAUSE)
    assert page.locator(".todo-list li").count() == 3
    toggles = page.locator(".todo-list li .toggle")
    highlight(page, toggles.nth(0)); toggles.nth(0).click()
    page.wait_for_timeout(STEP_PAUSE)
    highlight(page, toggles.nth(1)); toggles.nth(1).click()
    page.wait_for_timeout(STEP_PAUSE)
    page.get_by_text("1 item left").wait_for(timeout=8000)
    page.get_by_role("link", name="Completed").click()
    page.wait_for_timeout(STEP_PAUSE)
    assert page.locator(".todo-list li").count() == 2, "completed filter mismatch"
    status = [(tasks[0], True), (tasks[1], True), (tasks[2], False)]
    stats = write_sprint_sheet(wb, status)
    return (f"3 sprint tasks created, {stats['done']} completed — "
            f"summary written to Sprint_Summary")


# ---------------------------------------------------------------- runner

def run_case(pw, case, wb):
    tag = re.sub(r"[^a-z0-9]+", "-", case["name"].lower()).strip("-")
    shot = os.path.join(SHOT_DIR, f"{tag}.png")
    outcome = {"case": case["name"], "status": "FAIL", "notes": "",
               "screenshot": shot}
    browser = pw.chromium.launch(proxy=_proxy_or_none())
    context = browser.new_context(record_video_dir=VIDEO_DIR,
                                  record_video_size={"width": 1280, "height": 800},
                                  ignore_https_errors=True)
    before = set(os.listdir(VIDEO_DIR)) if os.path.isdir(VIDEO_DIR) else set()
    page = context.new_page()
    try:
        exp = case["expected"]
        if exp == "pass":
            outcome["notes"] = do_purchase(page, case)
        elif exp == "fail":
            outcome["notes"] = do_locked_out(page, case)
        elif exp == "report":
            outcome["notes"] = do_report(page, case, wb)
        elif exp == "sprint":
            outcome["notes"] = do_sprint(page, case, wb)
        else:
            raise ValueError(f"unknown expected outcome: {exp!r}")
        outcome["status"] = "PASS"
    except AssertionError as e:
        outcome["notes"] = f"Check failed: {e}"
    except Exception as e:
        outcome["notes"] = f"Unexpected: {type(e).__name__}: {str(e)[:120]}"
        traceback.print_exc()
    finally:
        try:
            page.screenshot(path=shot)
        except Exception:
            pass
        context.close()  # flushes the video file
        browser.close()
        for f in sorted(set(os.listdir(VIDEO_DIR)) - before):
            if f.endswith(".webm"):
                dst = os.path.join(VIDEO_DIR, f"{tag}.webm")
                if f != f"{tag}.webm" and not os.path.exists(dst):
                    os.rename(os.path.join(VIDEO_DIR, f), dst)
    return outcome


def main():
    os.makedirs(SHOT_DIR, exist_ok=True)
    os.makedirs(VIDEO_DIR, exist_ok=True)
    cases = load_cases()
    print(f"Loaded {len(cases)} case(s) from test_data.xlsx")
    wb = new_results_wb()
    ws = wb["Results"]
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    passed = 0
    with sync_playwright() as pw:
        for i, case in enumerate(cases, 1):
            print(f"[{i}/{len(cases)}] {case['name']} ... ", end="", flush=True)
            oc = run_case(pw, case, wb)
            ws.append([oc["case"], oc["status"], oc["notes"], stamp])
            r = ws.max_row
            cell = ws.cell(row=r, column=2)
            cell.fill = GREEN if oc["status"] == "PASS" else RED
            cell.font = GREEN_FONT if oc["status"] == "PASS" else RED_FONT
            cell.alignment = CENTER
            for c in (1, 3, 4):
                ws.cell(row=r, column=c).border = THIN
                ws.cell(row=r, column=c).alignment = LEFT if c != 4 else CENTER
            ws.cell(row=r, column=2).border = THIN
            if oc["status"] == "PASS":
                passed += 1
            print(f"{oc['status']} - {oc['notes'][:70]}")
    wb.save(RESULTS_XLSX)
    print(f"\nDone: {passed}/{len(cases)} PASS — report: {RESULTS_XLSX}")


if __name__ == "__main__":
    sys.exit(main())
